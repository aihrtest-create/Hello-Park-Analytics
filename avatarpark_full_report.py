import requests
import pandas as pd
import json
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =========================================
# КОНФИГУРАЦИЯ
# =========================================
GRAFANA_URL    = "http://10.0.110.11:3000"
API_TOKEN      = "YOUR_GRAFANA_API_TOKEN"
DATASOURCE_UID = "adg0kymjjbf28a"

OUTPUT_FILE = "/Users/dima/Downloads/avatarpark_report_2025_2026.xlsx"

# =========================================
# ПАРКИ
# =========================================
PARKS = [
    "AVIAPARK", "MEGA", "RIVIERA", "SELIGERSKAYA",
    "VLADIKAVKAZ", "SOCHI", "SAKHALIN", "BAKU",
    "DUBAI", "BOGOTA-NUESTRO", "Atyrau", "Kaspiysk",
]

# =========================================
# ПЕРИОДЫ (с поправкой -3ч для МСК)
# =========================================
PERIODS = [
    ("Янв 2025",  "2024-12-31T21:00:00Z", "2025-01-31T21:00:00Z"),
    ("Фев 2025",  "2025-01-31T21:00:00Z", "2025-02-28T21:00:00Z"),
    ("Мар 2025",  "2025-02-28T21:00:00Z", "2025-03-31T21:00:00Z"),
    ("Апр 2025",  "2025-03-31T21:00:00Z", "2025-04-30T21:00:00Z"),
    ("Май 2025",  "2025-04-30T21:00:00Z", "2025-05-31T21:00:00Z"),
    ("Июн 2025",  "2025-05-31T21:00:00Z", "2025-06-30T21:00:00Z"),
    ("Июл 2025",  "2025-06-30T21:00:00Z", "2025-07-31T21:00:00Z"),
    ("Авг 2025",  "2025-07-31T21:00:00Z", "2025-08-31T21:00:00Z"),
    ("Сен 2025",  "2025-08-31T21:00:00Z", "2025-09-30T21:00:00Z"),
    ("Окт 2025",  "2025-09-30T21:00:00Z", "2025-10-31T21:00:00Z"),
    ("Ноя 2025",  "2025-10-31T21:00:00Z", "2025-11-30T21:00:00Z"),
    ("Дек 2025",  "2025-11-30T21:00:00Z", "2025-12-31T21:00:00Z"),
    ("Янв 2026",  "2025-12-31T21:00:00Z", "2026-01-31T21:00:00Z"),
    ("Фев 2026",  "2026-01-31T21:00:00Z", "2026-02-16T21:00:00Z"),
]

# =========================================
# МАППИНГ zone -> название
# =========================================
ZONE_NAMES = {
    "hp-avatars-camp":    "Лагерь аватаров",
    "hp-boxgame":         "Кубики",
    "hp-bubble-mubble":   "Бабл-Мабл",
    "hp-disco":           "Диско",
    "hp-flashlights":     "Фонари",
    "hp-kidalki":         "Кидалки",
    "hp-living-figures":  "Ожившие фигуры",
    "hp-living-pictures": "Ожившие рисунки",
    "hp-painter":         "Художник",
    "hp-quest-portal":    "Квест-портал",
    "hp-trampoline":      "Батут",
}

HEADERS = {
    "Authorization": f"Bearer {API_TOKEN}",
    "Content-Type":  "application/json",
}

# =========================================
# ЗАПРОС К GRAFANA
# =========================================
def run_query(flux, d1, d2, label=""):
    payload = {
        "queries": [{
            "refId": "A",
            "datasource": {"type": "influxdb", "uid": DATASOURCE_UID},
            "query": flux,
        }],
        "from": d1,
        "to":   d2,
    }
    try:
        r = requests.post(
            f"{GRAFANA_URL}/api/ds/query",
            headers=HEADERS, json=payload, timeout=60
        )
    except Exception as e:
        print(f"    [Ошибка соединения] {label}: {e}")
        return {}
    if r.status_code != 200:
        print(f"    [HTTP {r.status_code}] {label}: {r.text[:200]}")
        return {}
    return r.json()

def parse_by_zone(data):
    """Парсит ответ -> {zone: count}"""
    out = {}
    try:
        frames = data.get("results", {}).get("A", {}).get("frames", [])
        for frame in frames:
            fields = frame.get("schema", {}).get("fields", [])
            values = frame.get("data", {}).get("values", [])
            labels = fields[0].get("labels", {}) if fields else {}
            zone   = labels.get("zone", "?")
            count  = values[0][0] if values and values[0] else 0
            if zone and zone != "?":
                out[zone] = count
    except Exception as e:
        print(f"    [Ошибка парсинга] {e}")
    return out

# =========================================
# FLUX-ЗАПРОСЫ
# =========================================
def q_scans(park, d1, d2):
    return f'''from(bucket: "Analytics_AvatarBD")
  |> range(start: {d1}, stop: {d2})
  |> filter(fn: (r) => r["_measurement"] == "PlayerEnteredInstallation"
                    and r["park"] == "{park}")
  |> group(columns: ["zone"])
  |> count(column: "_value")
  |> yield(name: "scans")'''

def q_achiev(park, d1, d2, name_filter):
    return f'''from(bucket: "Analytics_AvatarBD")
  |> range(start: {d1}, stop: {d2})
  |> filter(fn: (r) => r["_measurement"] == "AchievementUnlocked"
                    and r["park"] == "{park}")
  |> filter(fn: (r) => r["name"] =~ /{name_filter}/)
  |> group(columns: ["zone"])
  |> count(column: "_value")
  |> yield(name: "{name_filter}")'''

# =========================================
# СБОР ДАННЫХ
# =========================================
# Структура: data[park][period_label][metric][zone] = value
# metrics: scans, easy, wins, hard

all_data = {}
total_periods = len(PERIODS)
total_parks   = len(PARKS)

print(f"Начинаю сбор данных: {total_parks} парков x {total_periods} периодов")
print("=" * 60)

for pi, park in enumerate(PARKS):
    all_data[park] = {}
    for mi, (period_label, d1, d2) in enumerate(PERIODS):
        print(f"  [{pi+1}/{total_parks}] {park} | {period_label}...", end=" ", flush=True)
        sc = parse_by_zone(run_query(q_scans(park, d1, d2),           d1, d2, "scans"))
        ea = parse_by_zone(run_query(q_achiev(park, d1, d2, "TimeSpent"), d1, d2, "easy"))
        wi = parse_by_zone(run_query(q_achiev(park, d1, d2, "GameWin"),   d1, d2, "wins"))
        ha = parse_by_zone(run_query(q_achiev(park, d1, d2, "GameAction"),d1, d2, "hard"))
        all_data[park][period_label] = {
            "scans": sc, "easy": ea, "wins": wi, "hard": ha
        }
        total_sc = sum(sc.values())
        total_wi = sum(wi.values())
        print(f"сканов={total_sc}, побед={total_wi}")

print("\nДанные собраны. Формирую Excel...")

# =========================================
# СПИСОК ВСЕХ ЗОН
# =========================================
all_zones = sorted(ZONE_NAMES.keys())
period_labels = [p[0] for p in PERIODS]

# =========================================
# СТИЛИ
# =========================================
ORANGE     = PatternFill("solid", fgColor="FF6B00")
LIGHT_GRAY = PatternFill("solid", fgColor="F2F2F2")
DARK_GRAY  = PatternFill("solid", fgColor="D9D9D9")
BLUE_LIGHT = PatternFill("solid", fgColor="E8F0FE")
GREEN_LIGHT= PatternFill("solid", fgColor="E6F4EA")

HDR_FONT   = Font(bold=True, color="FFFFFF", size=10)
BOLD_FONT  = Font(bold=True, size=10)
REG_FONT   = Font(size=10)
CENTER     = Alignment(horizontal="center", vertical="center")
LEFT       = Alignment(horizontal="left",   vertical="center")

def style_header(cell):
    cell.fill      = ORANGE
    cell.font      = HDR_FONT
    cell.alignment = CENTER

def style_subheader(cell):
    cell.fill      = DARK_GRAY
    cell.font      = BOLD_FONT
    cell.alignment = CENTER

def style_total(cell):
    cell.fill = LIGHT_GRAY
    cell.font = BOLD_FONT

def col_width(ws, col_letter, width):
    ws.column_dimensions[col_letter].width = width

# =========================================
# EXCEL WRITER
# =========================================
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:

    # ─────────────────────────────────────
    # ЛИСТ 1: СВОДКА — итоги за весь период
    # ─────────────────────────────────────
    print("  Лист 1: Сводка...")
    rows = []
    for zone in all_zones:
        name = ZONE_NAMES.get(zone, zone)
        total_sc = total_ea = total_wi = total_ha = 0
        for park in PARKS:
            for pl in period_labels:
                d = all_data[park][pl]
                total_sc += d["scans"].get(zone, 0) or 0
                total_ea += d["easy"].get(zone, 0)  or 0
                total_wi += d["wins"].get(zone, 0)  or 0
                total_ha += d["hard"].get(zone, 0)  or 0
        conv = f"{total_wi/total_sc*100:.1f}%" if total_sc > 0 else "-"
        rows.append({
            "Игра":                  name,
            "Сканирований (всего)":  total_sc,
            "Лёгкие (TimeSpent)":    total_ea,
            "Победы (GameWin)":      total_wi,
            "Сложные (GameAction)":  total_ha,
            "Конверсия скан→победа": conv,
        })

    df_summary = pd.DataFrame(rows).sort_values("Сканирований (всего)", ascending=False)
    df_summary.to_excel(writer, sheet_name="1. Сводка", index=False)
    ws = writer.sheets["1. Сводка"]
    for cell in ws[1]:
        style_header(cell)
    widths = [22, 20, 20, 18, 20, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ─────────────────────────────────────
    # ЛИСТ 2: ПОПУЛЯРНОСТЬ (сканирования по месяцам x паркам)
    # ─────────────────────────────────────
    print("  Лист 2: Популярность по месяцам...")
    rows2 = []
    for zone in all_zones:
        row = {"Игра": ZONE_NAMES.get(zone, zone)}
        for park in PARKS:
            for pl in period_labels:
                val = all_data[park][pl]["scans"].get(zone, 0) or 0
                row[f"{park} | {pl}"] = val
        # Итого по игре
        row["ИТОГО"] = sum(
            all_data[p][pl]["scans"].get(zone, 0) or 0
            for p in PARKS for pl in period_labels
        )
        rows2.append(row)

    df2 = pd.DataFrame(rows2)
    df2 = df2.sort_values("ИТОГО", ascending=False)
    df2.to_excel(writer, sheet_name="2. Популярность", index=False)
    ws2 = writer.sheets["2. Популярность"]
    for cell in ws2[1]:
        style_header(cell)
    ws2.column_dimensions["A"].width = 22
    for i in range(2, len(df2.columns) + 1):
        ws2.column_dimensions[get_column_letter(i)].width = 14

    # ─────────────────────────────────────
    # ЛИСТ 3: ПОБЕДЫ (GameWin по месяцам x паркам)
    # ─────────────────────────────────────
    print("  Лист 3: Победы по месяцам...")
    rows3 = []
    for zone in all_zones:
        row = {"Игра": ZONE_NAMES.get(zone, zone)}
        for park in PARKS:
            for pl in period_labels:
                val = all_data[park][pl]["wins"].get(zone, 0) or 0
                row[f"{park} | {pl}"] = val
        row["ИТОГО"] = sum(
            all_data[p][pl]["wins"].get(zone, 0) or 0
            for p in PARKS for pl in period_labels
        )
        rows3.append(row)

    df3 = pd.DataFrame(rows3)
    df3 = df3.sort_values("ИТОГО", ascending=False)
    df3.to_excel(writer, sheet_name="3. Победы", index=False)
    ws3 = writer.sheets["3. Победы"]
    for cell in ws3[1]:
        style_header(cell)
    ws3.column_dimensions["A"].width = 22
    for i in range(2, len(df3.columns) + 1):
        ws3.column_dimensions[get_column_letter(i)].width = 14

    # ─────────────────────────────────────
    # ЛИСТ 4: ВОРОНКА по играм
    # ─────────────────────────────────────
    print("  Лист 4: Воронка...")
    rows4 = []
    for zone in all_zones:
        name = ZONE_NAMES.get(zone, zone)
        sc = ea = wi = ha = 0
        for park in PARKS:
            for pl in period_labels:
                d = all_data[park][pl]
                sc += d["scans"].get(zone, 0) or 0
                ea += d["easy"].get(zone, 0)  or 0
                wi += d["wins"].get(zone, 0)  or 0
                ha += d["hard"].get(zone, 0)  or 0
        rows4.append({
            "Игра":                       name,
            "1. Сканирований":            sc,
            "2. Лёгкие (TimeSpent)":      ea,
            "% скан→лёгкая":              f"{ea/sc*100:.1f}%" if sc > 0 else "-",
            "3. Победы (GameWin)":        wi,
            "% скан→победа":              f"{wi/sc*100:.1f}%" if sc > 0 else "-",
            "4. Сложные (GameAction)":    ha,
            "% скан→сложная":             f"{ha/sc*100:.1f}%" if sc > 0 else "-",
        })

    df4 = pd.DataFrame(rows4).sort_values("1. Сканирований", ascending=False)
    df4.to_excel(writer, sheet_name="4. Воронка", index=False)
    ws4 = writer.sheets["4. Воронка"]
    for cell in ws4[1]:
        style_header(cell)
    widths4 = [22, 16, 20, 16, 18, 16, 20, 16]
    for i, w in enumerate(widths4, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    # ─────────────────────────────────────
    # ЛИСТ 5: МАТРИЦА (трафик vs вовлечённость)
    # ─────────────────────────────────────
    print("  Лист 5: Матрица...")
    matrix_rows = []
    for zone in all_zones:
        name = ZONE_NAMES.get(zone, zone)
        sc = wi = ha = 0
        for park in PARKS:
            for pl in period_labels:
                d = all_data[park][pl]
                sc += d["scans"].get(zone, 0) or 0
                wi += d["wins"].get(zone, 0)  or 0
                ha += d["hard"].get(zone, 0)  or 0
        win_pct  = wi/sc*100 if sc > 0 else 0
        hard_pct = ha/sc*100 if sc > 0 else 0
        # Квадрант
        med_sc   = 1  # будет пересчитан ниже
        if sc > 0:
            category = "TBD"
        else:
            category = "Нет данных"
        matrix_rows.append({
            "Игра":               name,
            "Сканирований":       sc,
            "% побед":            round(win_pct, 1),
            "% сложных ачивок":   round(hard_pct, 1),
            "Категория":          category,
        })

    df5 = pd.DataFrame(matrix_rows)
    # Считаем медиану для разделения на квадранты
    med_scans = df5["Сканирований"].median()
    med_wins  = df5["% побед"].median()

    def get_category(row):
        if row["Сканирований"] == 0:
            return "Нет данных"
        high_traffic  = row["Сканирований"] >= med_scans
        high_engage   = row["% побед"]      >= med_wins
        if high_traffic and high_engage:
            return "Хит"
        elif high_traffic and not high_engage:
            return "Проходная (много идут, мало побеждают)"
        elif not high_traffic and high_engage:
            return "Скрытая жемчужина"
        else:
            return "Аутсайдер"

    df5["Категория"] = df5.apply(get_category, axis=1)
    df5 = df5.sort_values("Сканирований", ascending=False)
    df5.to_excel(writer, sheet_name="5. Матрица", index=False)
    ws5 = writer.sheets["5. Матрица"]
    for cell in ws5[1]:
        style_header(cell)

    # Цветовая маркировка по категории
    cat_colors = {
        "Хит":                                   "C6EFCE",
        "Проходная (много идут, мало побеждают)": "FFEB9C",
        "Скрытая жемчужина":                      "BDD7EE",
        "Аутсайдер":                              "FFC7CE",
        "Нет данных":                             "F2F2F2",
    }
    for row in ws5.iter_rows(min_row=2, max_row=ws5.max_row):
        cat_cell = row[4]  # колонка E = Категория
        color = cat_colors.get(cat_cell.value, "FFFFFF")
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=color)

    widths5 = [22, 16, 12, 18, 38]
    for i, w in enumerate(widths5, 1):
        ws5.column_dimensions[get_column_letter(i)].width = w

print(f"\nГотово! Файл сохранён: {OUTPUT_FILE}")
print("\nЛисты в файле:")
print("  1. Сводка          — итоги за весь период по всем паркам")
print("  2. Популярность    — сканирования по месяцам и паркам")
print("  3. Победы          — GameWin по месяцам и паркам")
print("  4. Воронка         — конверсия по каждому шагу")
print("  5. Матрица         — хиты / жемчужины / проходные / аутсайдеры")
