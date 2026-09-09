import requests
import pandas as pd
from backend.reports.common import GRAFANA_URL, DATASOURCE_UID, get_headers, get_time_boundaries, normalize_park_name
from openpyxl.styles import PatternFill, Font, Alignment

def generate_repeat_visits(date_val: str, period_type: str, selected_parks: list, output_path: str):
    start_date, stop_date, _ = get_time_boundaries(date_val, period_type)
    
    # 1. Шаг 1: Кол-во привязанных аватаров (AvatarLinkedInHome)
    flux_avatars = f"""
    from(bucket: \"Analytics_AvatarBD\")
      |> range(start: {start_date}, stop: {stop_date})
      |> filter(fn: (r) => r[\"_measurement\"] == \"AvatarLinkedInHome\")
      |> filter(fn: (r) => exists r[\"id\"])
      |> group(columns: [\"park\", \"id\"])
      |> limit(n: 1)
      |> group(columns: [\"park\"])
      |> count(column: \"_value\")
      |> yield(name: \"avatars\")
    """
    
    # 2. Шаг 2 и 3: Повторные посещения (userBindRfid c дедупликацией 12h)
    flux_visits = f"""
    DedupWindow = 12h

    binds = from(bucket: \"Analytics_AvatarBD\")
      |> range(start: {start_date}, stop: {stop_date})
      |> filter(fn: (r) => r._measurement == \"userBindRfid\")
      |> filter(fn: (r) => r._field == \"user_id\")
      |> duplicate(column: \"_value\", as: \"id\")
      |> map(fn: (r) => ({{ r with _value: 1 }}))
      |> keep(columns: [\"_time\", \"_value\", \"id\", \"park\"])
      |> group(columns: [\"park\", \"id\"])
      |> sort(columns: [\"_time\"])

    firstBinds = binds |> first()

    laterBinds = binds
      |> elapsed(unit: 1ns)
      |> filter(fn: (r) => r.elapsed >= int(v: DedupWindow))

    union(tables: [firstBinds, laterBinds])
      |> group(columns: [\"park\", \"id\"])
      |> count(column: \"_value\")
      |> group(columns: [\"park\", \"_value\"])
      |> count(column: \"id\")
      |> yield(name: \"visits\")
    """
    
    # Выполняем запрос аватаров
    avatars_data = {}
    try:
        p_av = {"queries": [{"refId": "A", "datasource": {"type": "influxdb", "uid": DATASOURCE_UID}, "query": flux_avatars}], "from": "0", "to": "9999999999999"}
        r_av = requests.post(f"{GRAFANA_URL}/api/ds/query", headers=get_headers(), json=p_av, timeout=45)
        r_av.raise_for_status()
        for frame in r_av.json().get("results", {}).get("A", {}).get("frames", []):
            park_name = "Unknown"
            for field in frame.get("schema", {}).get("fields", []):
                if "park" in field.get("labels", {}):
                    park_name = field["labels"]["park"]
                    break
            park_name = normalize_park_name(park_name)
            vals = frame.get("data", {}).get("values", [])
            if len(vals) > 0 and len(vals[0]) > 0:
                val = vals[0][0] or 0
                avatars_data[park_name] = avatars_data.get(park_name, 0) + val
    except Exception as e:
        print(f"Error querying AvatarLinkedInHome: {e}")

    # Выполняем запрос посещений
    visits_data = {}
    try:
        p_vis = {"queries": [{"refId": "A", "datasource": {"type": "influxdb", "uid": DATASOURCE_UID}, "query": flux_visits}], "from": "0", "to": "9999999999999"}
        r_vis = requests.post(f"{GRAFANA_URL}/api/ds/query", headers=get_headers(), json=p_vis, timeout=60)
        r_vis.raise_for_status()
        for frame in r_vis.json().get("results", {}).get("A", {}).get("frames", []):
            labels = {}
            for field in frame.get("schema", {}).get("fields", []):
                labels.update(field.get("labels", {}))
            park_name = normalize_park_name(labels.get("park", "Unknown"))
            try:
                visit_count = int(labels.get("_value", 0))
            except (ValueError, TypeError):
                visit_count = 0
            vals = frame.get("data", {}).get("values", [])
            user_count = vals[0][0] if vals and len(vals[0]) > 0 and vals[0][0] is not None else 0
            
            if park_name not in visits_data:
                visits_data[park_name] = {}
            visits_data[park_name][visit_count] = visits_data[park_name].get(visit_count, 0) + user_count
    except Exception as e:
        print(f"Error querying userBindRfid: {e}")

    # Формируем список парков
    all_parks = sorted([
        p for p in set(avatars_data.keys()).union(set(visits_data.keys()))
        if p not in ["OFFICE", "QA-TTUZOV", "Unknown"] and p in selected_parks
    ])

    rows = []
    for park in all_parks:
        av = avatars_data.get(park, 0)
        counts = visits_data.get(park, {})
        v2 = sum(cnt for v, cnt in counts.items() if v >= 2)
        v3 = sum(cnt for v, cnt in counts.items() if v >= 3)
        pct2 = round(v2 / av * 100, 2) if av > 0 else 0.0
        pct3 = round(v3 / av * 100, 2) if av > 0 else 0.0
        rows.append({
            "Парк": park,
            "Кол-во созданных": av,
            "Пришли 2 раз": v2,
            "% от Созданных_2": pct2,
            "Пришли 3 раз": v3,
            "% от Созданных_3": pct3
        })

    # Итоговые значения
    total_av = sum([r["Кол-во созданных"] for r in rows])
    total_v2 = sum([r["Пришли 2 раз"] for r in rows])
    total_pct2 = round(total_v2 / total_av * 100, 2) if total_av > 0 else 0.0
    total_v3 = sum([r["Пришли 3 раз"] for r in rows])
    total_pct3 = round(total_v3 / total_av * 100, 2) if total_av > 0 else 0.0

    rows.append({
        "Парк": "ИТОГО",
        "Кол-во созданных": total_av,
        "Пришли 2 раз": total_v2,
        "% от Созданных_2": total_pct2,
        "Пришли 3 раз": total_v3,
        "% от Созданных_3": total_pct3
    })

    df = pd.DataFrame(rows)
    df_data = df.iloc[:-1].copy().sort_values("Парк")
    df_final = pd.concat([df_data, df.iloc[[-1]]], ignore_index=True)

    # Переименовываем колонки для финальной таблицы
    df_final.columns = [
        "Парк",
        "Кол-во созданных",
        "Пришли 2 раз",
        "% от Созданных",
        "Пришли 3 раз",
        "% от Созданных"
    ]

    sheet_title = f"Повторные визиты {date_val}"[:31]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name=sheet_title, index=False)
        ws = writer.sheets[sheet_title]
        
        orange = PatternFill("solid", fgColor="FF6B00")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center = Alignment(horizontal="center", vertical="center")
        gray = PatternFill("solid", fgColor="E0E0E0")
        bold_font = Font(bold=True)
        green = PatternFill("solid", fgColor="D5F5E3")
        
        # Шапка
        for col in range(1, len(df_final.columns) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = orange
            cell.font = header_font
            cell.alignment = center
            
        last_row = ws.max_row
        
        # Строка ИТОГО
        for col in range(1, len(df_final.columns) + 1):
            cell = ws.cell(row=last_row, column=col)
            cell.fill = gray
            cell.font = bold_font
            
        # Форматирование колонок с процентами (D и F) и числами
        for row in range(2, last_row + 1):
            # Проценты для 2-го визита (Колонка D / 4)
            cell_pct2 = ws.cell(row=row, column=4)
            cell_pct2.fill = green
            cell_pct2.number_format = '0.00"%"'
            cell_pct2.alignment = Alignment(horizontal="right")
            
            # Проценты для 3-го визита (Колонка F / 6)
            cell_pct3 = ws.cell(row=row, column=6)
            cell_pct3.fill = green
            cell_pct3.number_format = '0.00"%"'
            cell_pct3.alignment = Alignment(horizontal="right")

            # Числа (B, C, E)
            for c_idx in [2, 3, 5]:
                ws.cell(row=row, column=c_idx).number_format = '#,##0'
                ws.cell(row=row, column=c_idx).alignment = Alignment(horizontal="right")
            
        # Ширина колонок
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 18

    return output_path
