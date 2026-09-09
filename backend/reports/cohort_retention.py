import calendar
from datetime import datetime, timezone
import pandas as pd
import requests
from openpyxl.styles import PatternFill, Font, Alignment
from backend.reports.common import GRAFANA_URL, DATASOURCE_UID, get_headers, normalize_park_name, PARK_ORDER

MONTH_NAMES_RU = [
    "", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
]

def get_month_boundaries(year: int, month: int):
    """Вычисляет границы календарного месяца с учётом часового пояса MSK (UTC+3)"""
    _, last_day = calendar.monthrange(year, month)
    prev_month = 12 if month == 1 else month - 1
    prev_year = year - 1 if month == 1 else year
    _, prev_last_day = calendar.monthrange(prev_year, prev_month)
    
    start_str = f"{prev_year:04d}-{prev_month:02d}-{prev_last_day:02d}T21:00:00Z"
    stop_str = f"{year:04d}-{month:02d}-{last_day:02d}T21:00:00Z"
    return start_str, stop_str

def parse_base_year_month(date_val: str):
    """Извлекает год и месяц из любого формата даты (YYYY-MM, YYYY-MM-DD, диапазон)"""
    clean = str(date_val).strip()
    try:
        parts = clean.split("-")
        y = int(parts[0])
        m = int(parts[1])
        return y, m
    except Exception:
        now = datetime.now()
        return now.year, now.month

def generate_cohort_retention(date_val: str, period_type: str, selected_parks: list, output_path: str):
    m0_year, m0_month = parse_base_year_month(date_val)
    now = datetime.now(timezone.utc)
    
    # Формируем цепочку из 5 месяцев (M0, M+1, M+2, M+3, M+4)
    months_info = []
    curr_y, curr_m = m0_year, m0_month
    
    for i in range(5):
        s_str, e_str = get_month_boundaries(curr_y, curr_m)
        s_dt = datetime.fromisoformat(s_str.replace("Z", "+00:00"))
        e_dt = datetime.fromisoformat(e_str.replace("Z", "+00:00"))
        
        if now < s_dt:
            status = "future"
        elif s_dt <= now <= e_dt:
            status = "current"
        else:
            status = "completed"
            
        month_label = f"{MONTH_NAMES_RU[curr_m]} {curr_y}"
        if status == "current":
            month_label += "*"
            
        months_info.append({
            "step": i,
            "label": month_label,
            "year": curr_y,
            "month": curr_m,
            "start_str": s_str,
            "stop_str": e_str,
            "s_ts": int(s_dt.timestamp() * 1000),
            "e_ts": int(e_dt.timestamp() * 1000),
            "status": status
        })
        
        if curr_m == 12:
            curr_m = 1
            curr_y += 1
        else:
            curr_m += 1

    m0_start = months_info[0]["start_str"]
    m0_stop = months_info[0]["stop_str"]
    
    # Дата окончания наблюдения - минимум из конца M4 и текущего момента
    obs_stop_dt = min(now, datetime.fromisoformat(months_info[-1]["stop_str"].replace("Z", "+00:00")))
    obs_stop = obs_stop_dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    
    # 1. Запрашиваем когорту M0: уникальные аватары по паркам
    flux_cohort = f"""
    from(bucket: "Analytics_AvatarBD")
      |> range(start: {m0_start}, stop: {m0_stop})
      |> filter(fn: (r) => r._measurement == "AvatarLinkedInHome" and exists r.id)
      |> group(columns: ["park"])
      |> unique(column: "id")
      |> keep(columns: ["park", "id"])
    """
    
    cohort_by_park = {}
    try:
        p1 = {
            "queries": [{"refId": "A", "datasource": {"type": "influxdb", "uid": DATASOURCE_UID}, "query": flux_cohort, "maxDataPoints": 50000}],
            "from": "0", "to": "9999999999999"
        }
        r1 = requests.post(f"{GRAFANA_URL}/api/ds/query", headers=get_headers(), json=p1, timeout=90)
        r1.raise_for_status()
        frames1 = r1.json().get("results", {}).get("A", {}).get("frames", [])
        for f in frames1:
            labels = {}
            for fld in f.get("schema", {}).get("fields", []):
                if fld.get("labels"):
                    labels.update(fld["labels"])
            p = normalize_park_name(labels.get("park", "Unknown"))
            vals = f.get("data", {}).get("values", [])
            if p not in cohort_by_park:
                cohort_by_park[p] = set()
            if vals and len(vals) > 0:
                cohort_by_park[p].update(vals[0])
    except Exception as e:
        print(f"Error querying cohort avatars: {e}", flush=True)

    # 2. Запрашиваем все визиты с дедупликацией 12h за общий интервал наблюдения
    visits_by_park = {}
    flux_visits = f"""
    DedupWindow = 12h

    binds = from(bucket: "Analytics_AvatarBD")
      |> range(start: {m0_start}, stop: {obs_stop})
      |> filter(fn: (r) => r._measurement == "userBindRfid" and r._field == "user_id")
      |> keep(columns: ["_time", "_value", "park"])
      |> group(columns: ["park", "_value"])
      |> sort(columns: ["_time"])

    firstBinds = binds |> first()

    laterBinds = binds
      |> elapsed(unit: 1ns)
      |> filter(fn: (r) => r.elapsed >= int(v: DedupWindow))

    union(tables: [firstBinds, laterBinds])
      |> keep(columns: ["_time", "_value", "park"])
      |> group(columns: ["park"])
    """
    
    try:
        p2 = {
            "queries": [{"refId": "A", "datasource": {"type": "influxdb", "uid": DATASOURCE_UID}, "query": flux_visits, "maxDataPoints": 100000}],
            "from": "0", "to": "9999999999999"
        }
        r2 = requests.post(f"{GRAFANA_URL}/api/ds/query", headers=get_headers(), json=p2, timeout=90)
        r2.raise_for_status()
        frames2 = r2.json().get("results", {}).get("A", {}).get("frames", [])
        for f in frames2:
            labels = {}
            for fld in f.get("schema", {}).get("fields", []):
                if fld.get("labels"):
                    labels.update(fld["labels"])
            p = normalize_park_name(labels.get("park", "Unknown"))
            vals = f.get("data", {}).get("values", [])
            if len(vals) >= 2:
                times = vals[0]
                uids = vals[1]
                if p not in visits_by_park:
                    visits_by_park[p] = {}
                for t, u in zip(times, uids):
                    if u not in visits_by_park[p]:
                        visits_by_park[p][u] = []
                    visits_by_park[p][u].append(t)
    except Exception as e:
        print(f"Error querying userBindRfid: {e}", flush=True)

    # Гарантируем включение всех выбранных парков
    clean_selected = [
        normalize_park_name(p) for p in selected_parks
        if p not in ["OFFICE", "QA-TTUZOV", "Unknown"]
    ]
    all_parks = sorted(list(set(clean_selected)))
    if not all_parks:
        all_parks = sorted([
            p for p in set(cohort_by_park.keys()).union(set(visits_by_park.keys()))
            if p not in ["OFFICE", "QA-TTUZOV", "Unknown"]
        ])

    m0_label = months_info[0]["label"]
    rows = []
    
    for park in all_parks:
        cohort_users = cohort_by_park.get(park, set())
        base_count = len(cohort_users)
        park_visits = visits_by_park.get(park, {})
        
        row = {
            "Парк": park,
            f"Когорта ({m0_label})": base_count
        }
        
        returned_any = set()
        
        for m in months_info:
            step = m["step"]
            lbl = m["label"]
            status = m["status"]
            s_ts = m["s_ts"]
            e_ts = m["e_ts"]
            
            col_name_abs = f"M0 ({lbl})" if step == 0 else f"+{step}M ({lbl})"
            col_name_pct = f"%_{step}"
            
            if status == "future":
                row[col_name_abs] = "—"
                row[col_name_pct] = "—"
                continue
                
            returned_month = set()
            for u in cohort_users:
                u_times = park_visits.get(u, [])
                m_visits = [t for t in u_times if s_ts <= t < e_ts]
                
                if step == 0:
                    # В базовом месяце повторный визит = >= 2 посещений
                    if len(m_visits) >= 2:
                        returned_month.add(u)
                        returned_any.add(u)
                else:
                    # В последующих месяцах = хотя бы 1 посещение
                    if len(m_visits) >= 1:
                        returned_month.add(u)
                        returned_any.add(u)
                        
            cnt = len(returned_month)
            pct = round(cnt / base_count * 100, 2) if base_count > 0 else 0.0
            row[col_name_abs] = cnt
            row[col_name_pct] = pct
            
        tot_cnt = len(returned_any)
        tot_pct = round(tot_cnt / base_count * 100, 2) if base_count > 0 else 0.0
        row["Всего вернулись (LTV)"] = tot_cnt
        row["%_ltv"] = tot_pct
        rows.append(row)

    # Строка ИТОГО
    total_cohort = sum([r[f"Когорта ({m0_label})"] for r in rows])
    row_total = {
        "Парк": "ИТОГО",
        f"Когорта ({m0_label})": total_cohort
    }
    
    for m in months_info:
        step = m["step"]
        lbl = m["label"]
        status = m["status"]
        
        col_name_abs = f"M0 ({lbl})" if step == 0 else f"+{step}M ({lbl})"
        col_name_pct = f"%_{step}"
        
        if status == "future":
            row_total[col_name_abs] = "—"
            row_total[col_name_pct] = "—"
        else:
            tot_step_abs = sum([r[col_name_abs] for r in rows if isinstance(r[col_name_abs], (int, float))])
            tot_step_pct = round(tot_step_abs / total_cohort * 100, 2) if total_cohort > 0 else 0.0
            row_total[col_name_abs] = tot_step_abs
            row_total[col_name_pct] = tot_step_pct

    tot_ltv_abs = sum([r["Всего вернулись (LTV)"] for r in rows if isinstance(r["Всего вернулись (LTV)"], (int, float))])
    tot_ltv_pct = round(tot_ltv_abs / total_cohort * 100, 2) if total_cohort > 0 else 0.0
    row_total["Всего вернулись (LTV)"] = tot_ltv_abs
    row_total["%_ltv"] = tot_ltv_pct
    
    rows.append(row_total)
    df_final = pd.DataFrame(rows)

    # Формируем красивые заголовки колонок (заменяем служебные %_step на "% от Когорты")
    final_cols = []
    pct_col_indices = []
    
    for idx, col in enumerate(df_final.columns):
        if col.startswith("%_"):
            final_cols.append("% от Когорты")
            pct_col_indices.append(idx + 1)
        else:
            final_cols.append(col)
            
    df_final.columns = final_cols

    sheet_title = f"Когорта {m0_label}"[:31]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name=sheet_title, index=False)
        ws = writer.sheets[sheet_title]
        
        orange = PatternFill("solid", fgColor="FF6B00")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        gray = PatternFill("solid", fgColor="E0E0E0")
        bold_font = Font(bold=True)
        green = PatternFill("solid", fgColor="D5F5E3")
        
        # Оформление шапки
        for col_idx in range(1, len(df_final.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = orange
            cell.font = header_font
            cell.alignment = center
            
        last_row = ws.max_row
        
        # Оформление строки ИТОГО
        for col_idx in range(1, len(df_final.columns) + 1):
            cell = ws.cell(row=last_row, column=col_idx)
            cell.fill = gray
            cell.font = bold_font
            
        # Форматирование данных
        for r_idx in range(2, last_row + 1):
            for c_idx in range(1, len(df_final.columns) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                
                # Колонки с процентами
                if c_idx in pct_col_indices:
                    cell.fill = green
                    if cell.value == "—":
                        cell.alignment = Alignment(horizontal="center")
                    else:
                        cell.number_format = '0.00"%"'
                        cell.alignment = Alignment(horizontal="right")
                elif c_idx == 1:
                    # Колонка "Парк"
                    cell.alignment = Alignment(horizontal="left")
                else:
                    # Числовые колонки
                    if cell.value == "—":
                        cell.alignment = Alignment(horizontal="center")
                    else:
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="right")

        # Настройка ширины колонок
        ws.column_dimensions["A"].width = 22
        for col_letter in ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]:
            ws.column_dimensions[col_letter].width = 17

    return output_path
